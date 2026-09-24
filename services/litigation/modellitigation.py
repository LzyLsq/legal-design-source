import sys
from pathlib import Path

sys.path.insert(0, str(next(p for p in Path(__file__).resolve().parents if (p / "shared").is_dir())))

from shared import config, web  # noqa: E402

import os
from flask import request, render_template, send_file, jsonify
from markupsafe import Markup
import markdown
from docx import Document
from docx.oxml.ns import qn
from bs4 import BeautifulSoup
from io import BytesIO

app = web.create_flask_app(__name__)

# 自定义过滤器，用于提取文件名
@app.template_filter('basename')
def basename_filter(path):
    return os.path.basename(path)

# 文件的根目录
docs_root = str(config.path_from(__file__, "laws"))
all_files = []

# 初始化文件列表
for root, dirs, files in os.walk(docs_root):
    for file in files:
        if file.endswith('.md'):
            all_files.append(os.path.join(root, file))

@app.route('/')
def index():
    return render_template('index.html', all_files=all_files)

def _resolve_doc(filename):
    """把用户传入的文件名限制在 laws 目录内，拒绝路径穿越。"""
    if not filename:
        return None
    return config.safe_path(docs_root, filename)


@app.route('/show_md', methods=['GET'])
def show_md():
    filename = request.args.get('filename')
    safe = _resolve_doc(filename)
    if safe is None:
        return "Invalid filename", 400
    content = md2html(safe)
    md_content = safe.read_text(encoding='utf-8')
    return render_template('show_md.html', content=content, md_content=md_content, filename=filename)

@app.route('/download_word', methods=['GET'])
def download_word():
    filename = request.args.get('filename')
    safe = _resolve_doc(filename)
    if safe is None or not safe.exists():
        return "File not found", 404
    return send_file(str(safe), as_attachment=True)

@app.route('/save_md', methods=['POST'])
def save_md():
    data = request.form
    custom_filename = data.get('custom_filename')
    content = data.get('content')

    if not custom_filename:
        return jsonify({'status': 'error', 'message': 'Invalid filename'}), 400
    custom_filename = os.path.basename(custom_filename)
    if not custom_filename.endswith('.md'):
        custom_filename += '.md'

    tem_dir = str(config.path_from(__file__, 'mymodel'))
    os.makedirs(tem_dir, exist_ok=True)
    tem_path = os.path.join(tem_dir, custom_filename)

    try:
        with open(tem_path, 'w', encoding='utf-8') as f:
            f.write(content)
        return jsonify({'status': 'success', 'message': 'File saved successfully'})
    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)}), 500

@app.route('/delete_temp_file', methods=['POST'])
def delete_temp_file():
    data = request.form
    filename = data.get('filename')
    try:
        safe = _resolve_doc(filename)
        if safe is not None and safe.exists():
            os.remove(safe)
        return 'File deleted successfully'
    except Exception as e:
        return str(e)

@app.route('/save_as_word', methods=['POST'])
def save_as_word():
    data = request.form
    custom_filename = data.get('custom_filename')
    content = data.get('content')

    # 创建 Word 文档
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # 将 Markdown 内容转换为 HTML 并解析
    html_content = markdown.markdown(content)
    soup = BeautifulSoup(html_content, 'html.parser')
    elements = soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'ul', 'ol', 'li', 'strong', 'em'])

    for element in elements:
        if element.name.startswith('h'):
            level = int(element.name[1])
            heading = doc.add_heading(level=level)
            heading.add_run(element.get_text()).bold = True
        elif element.name == 'p':
            para = doc.add_paragraph(element.get_text())
        elif element.name in ['ul', 'ol']:
            for li in element.find_all('li'):
                para = doc.add_paragraph()
                run = para.add_run(li.get_text())
                if element.name == 'ul':
                    para.style = 'List Bullet'
                else:
                    para.style = 'List Number'
        elif element.name == 'strong':
            para = doc.add_paragraph()
            run = para.add_run(element.get_text())
            run.bold = True
        elif element.name == 'em':
            para = doc.add_paragraph()
            run = para.add_run(element.get_text())
            run.italic = True

    # # 将 Word 文档保存到内存中
    docx_io = BytesIO()
    doc.save(docx_io)
    docx_io.seek(0)  # 确保从头开始读取

    # 保存到模板目录
    tem_dir = str(config.path_from(__file__, 'mymodel'))
    os.makedirs(tem_dir, exist_ok=True)
    tem_path = os.path.join(tem_dir, f'{custom_filename}.docx')

    try:
        with open(tem_path, 'wb') as f:
            f.write(docx_io.getvalue())

        return jsonify({'status': 'success', 'message': 'File saved successfully'})
    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)}), 500

def md2html(filename):
    with open(filename, 'r', encoding='utf-8') as f:
        text = f.read()
    html = markdown.markdown(text)
    return Markup(html)

if __name__ == '__main__':
    web.run_flask(app, 5025)